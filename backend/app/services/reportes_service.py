"""Generación de los reportes descargables en Excel del módulo de Remisiones.

Los encabezados y el orden de columnas son los que pidió el cliente (ver
``backend/app/data_import/formatos/reporte-remisiones-*-programa_nuevo.xlsx``,
columna E de cada archivo) — no se pueden reordenar ni renombrar sin que el
cliente lo confirme. Ver NOTAS_RETOMAR.md / informe de entrega para los
supuestos de mapeo de cada columna "ID".
"""

from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select

from app.models.activo import Activo
from app.models.catalogos import FormaPago, Producto, Proyecto, TipoServicio, Tramo, UnidadMedida
from app.models.facturacion import Factura, RemisionFacturacion
from app.models.orden_servicio import OrdenServicio
from app.models.remision import Remision, RemisionProducto
from app.models.tercero import Tercero
from app.models.ubicacion import Municipio
from app.models.usuario import Usuario

COLUMNAS_GENERAL = [
    "ID",
    "Numero de remision",
    "Fecha",
    "Forma de pago",
    "ODS",
    "ID del cliente",
    "Cliente",
    "Proyecto",
    "Tramo",
    "Origen",
    "Destino",
    "Kilometros",
    "Vehiculos",
    "Cama Baja",
    "Hora llegada",
    "Hola salida",
    "ID del conductor",
    "Conductor",
    "Tipo de transporte",
    "Observaciones",
    "Id Producto",
    "Producto",
    "Cantidad",
    "Unidad de medida",
    "Valor Material",
    "Valor Flete",
    "Numero de viajes",
    "Valor de comision",
    "Valor de costo",
    "Valor cama baja",
    "Remision AGRESUR",
    "ID del usuario",
    "Usuario",
    "Estado",
]

COLUMNAS_FACTURADAS = [
    "Numero de remision",
    "Fecha",
    "Forma de pago",
    "ODS",
    "ID cliente",
    "Cliente",
    "Proyecto",
    "Tramo",
    "Origen",
    "Destino",
    "Kilometros",
    "Vehiculos",
    "Cama Baja",
    "Hora llegada",
    "Hola salida",
    "Conductor",
    "Tipo de transporte",
    "Usuario",
    "ID usuario",
    "ID factura",
    "Usuario Factura",
    "ID usuario factura",
    "Número de factura",
    "Fecha de factura",
    "Valor flete",
    "Valor Camabaja",
    "Subtotal",
    "Valor Material",
    "Total",
    "Valor facturado",
    "Valor pendiente",
    "Estado",
]

# Columnas que deben quedar como texto aunque el valor sea numérico (para no
# perder ceros a la izquierda ni que Excel las reinterprete).
COLUMNAS_TEXTO = {
    "ID del cliente",
    "ID cliente",
    "ID del conductor",
}

COLUMNAS_FECHA = {"Fecha", "Fecha de factura"}


def _aplicar_filtros(
    statement,
    *,
    desde: date | None,
    hasta: date | None,
    cliente_id: int | None,
    proyecto_id: int | None,
    vehiculo_id: int | None,
    camabaja_id: int | None,
    estado_facturacion: str | None = None,
):
    if desde:
        statement = statement.where(Remision.fecha >= desde)
    if hasta:
        statement = statement.where(Remision.fecha <= hasta)
    if cliente_id:
        statement = statement.where(Remision.cliente_id == cliente_id)
    if proyecto_id:
        statement = statement.where(Remision.proyecto_id == proyecto_id)
    if vehiculo_id:
        statement = statement.where(Remision.vehiculo_id == vehiculo_id)
    if camabaja_id:
        statement = statement.where(Remision.camabaja_id == camabaja_id)
    if estado_facturacion:
        statement = statement.where(Remision.estado_facturacion == estado_facturacion)
    return statement


class _Catalogos:
    """Cachés simples por id para no repetir `session.get` por cada fila."""

    def __init__(self, session: Session):
        self.session = session
        self._cache: dict[tuple[type, int], Any] = {}

    def get(self, modelo: type, id_: int | None):
        if id_ is None:
            return None
        clave = (modelo, id_)
        if clave not in self._cache:
            self._cache[clave] = self.session.get(modelo, id_)
        return self._cache[clave]


def _datos_comunes(cat: _Catalogos, remision: Remision) -> dict[str, Any]:
    cliente = cat.get(Tercero, remision.cliente_id)
    conductor = cat.get(Tercero, remision.conductor_id)
    origen = cat.get(Municipio, remision.origen_municipio_id)
    destino = cat.get(Municipio, remision.destino_municipio_id)
    proyecto = cat.get(Proyecto, remision.proyecto_id)
    tramo = cat.get(Tramo, remision.tramo_id)
    forma_pago = cat.get(FormaPago, remision.forma_pago_id)
    tipo_servicio = cat.get(TipoServicio, remision.tipo_servicio_id)
    vehiculo = cat.get(Activo, remision.vehiculo_id)
    camabaja = cat.get(Activo, remision.camabaja_id)
    usuario_creador = cat.get(Usuario, remision.usuario_creador_id)
    orden_servicio = cat.get(OrdenServicio, remision.orden_servicio_id)

    return {
        "id": remision.id,
        "numero_remision": remision.numero_remision,
        "fecha": remision.fecha,
        "forma_pago": forma_pago.nombre if forma_pago else None,
        "ods": orden_servicio.numero if orden_servicio else None,
        "cliente_identificacion": cliente.identificacion if cliente else None,
        "cliente_nombre": cliente.nombre if cliente else None,
        "proyecto": proyecto.nombre if proyecto else None,
        "tramo": tramo.nombre if tramo else None,
        "origen": origen.nombre if origen else None,
        "destino": destino.nombre if destino else None,
        "kilometros": remision.kilometros,
        "vehiculo_placa": vehiculo.placa if vehiculo else None,
        "camabaja_placa": camabaja.placa if camabaja else None,
        "hora_llegada": remision.hora_llegada,
        "hora_salida": remision.hora_salida,
        "conductor_identificacion": conductor.identificacion if conductor else None,
        "conductor_nombre": conductor.nombre if conductor else None,
        "tipo_servicio": tipo_servicio.nombre if tipo_servicio else None,
        "observaciones": remision.observaciones,
        "remision_agresur": remision.remision_agresur,
        "usuario_creador_id": usuario_creador.id if usuario_creador else None,
        "usuario_creador_nombre": usuario_creador.nombre if usuario_creador else None,
        "estado_facturacion": remision.estado_facturacion,
    }


def _totales_remision(session: Session, remision_id: int) -> dict[str, Decimal]:
    productos = session.exec(
        select(RemisionProducto).where(RemisionProducto.remision_id == remision_id)
    ).all()
    valor_material = sum((p.valor_material for p in productos), Decimal(0))
    valor_flete = sum((p.valor_flete for p in productos), Decimal(0))
    valor_camabaja = sum((p.valor_camabaja for p in productos), Decimal(0))
    subtotal = valor_flete + valor_camabaja
    return {
        "valor_material": valor_material,
        "valor_flete": valor_flete,
        "valor_camabaja": valor_camabaja,
        "subtotal": subtotal,
        "total": subtotal + valor_material,
    }


def _escribir_xlsx(filas: list[dict[str, Any]], columnas: list[str]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for col_idx, nombre in enumerate(columnas, start=1):
        ws.cell(row=1, column=col_idx, value=nombre)

    for row_idx, fila in enumerate(filas, start=2):
        for col_idx, nombre in enumerate(columnas, start=1):
            valor = fila.get(nombre)
            celda = ws.cell(row=row_idx, column=col_idx)
            if nombre in COLUMNAS_TEXTO:
                celda.value = str(valor) if valor is not None else None
                celda.number_format = "@"
            elif nombre in COLUMNAS_FECHA:
                celda.value = valor
                if valor is not None:
                    celda.number_format = "yyyy-mm-dd"
            elif isinstance(valor, Decimal):
                celda.value = float(valor)
            else:
                celda.value = valor

    for col_idx in range(1, len(columnas) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generar_reporte_general(
    session: Session,
    *,
    desde: date | None,
    hasta: date | None,
    cliente_id: int | None,
    proyecto_id: int | None,
    vehiculo_id: int | None,
    camabaja_id: int | None,
    estado_facturacion: str | None = None,
) -> BytesIO:
    statement = _aplicar_filtros(
        select(Remision),
        desde=desde,
        hasta=hasta,
        cliente_id=cliente_id,
        proyecto_id=proyecto_id,
        vehiculo_id=vehiculo_id,
        camabaja_id=camabaja_id,
        estado_facturacion=estado_facturacion,
    )
    remisiones = session.exec(statement.order_by(Remision.fecha, Remision.id)).all()
    cat = _Catalogos(session)

    filas: list[dict[str, Any]] = []
    for remision in remisiones:
        comunes = _datos_comunes(cat, remision)
        lineas = session.exec(
            select(RemisionProducto)
            .where(RemisionProducto.remision_id == remision.id)
            .order_by(RemisionProducto.id)
        ).all()

        base_fila = {
            "ID": comunes["id"],
            "Numero de remision": comunes["numero_remision"],
            "Fecha": comunes["fecha"],
            "Forma de pago": comunes["forma_pago"],
            "ODS": comunes["ods"],
            "ID del cliente": comunes["cliente_identificacion"],
            "Cliente": comunes["cliente_nombre"],
            "Proyecto": comunes["proyecto"],
            "Tramo": comunes["tramo"],
            "Origen": comunes["origen"],
            "Destino": comunes["destino"],
            "Kilometros": comunes["kilometros"],
            "Vehiculos": comunes["vehiculo_placa"],
            "Cama Baja": comunes["camabaja_placa"],
            "Hora llegada": comunes["hora_llegada"],
            "Hola salida": comunes["hora_salida"],
            "ID del conductor": comunes["conductor_identificacion"],
            "Conductor": comunes["conductor_nombre"],
            "Tipo de transporte": comunes["tipo_servicio"],
            "Observaciones": comunes["observaciones"],
            "Remision AGRESUR": comunes["remision_agresur"],
            "ID del usuario": comunes["usuario_creador_id"],
            "Usuario": comunes["usuario_creador_nombre"],
            "Estado": comunes["estado_facturacion"],
        }

        if not lineas:
            filas.append(dict(base_fila))
            continue

        for linea in lineas:
            producto = cat.get(Producto, linea.producto_id)
            unidad = cat.get(UnidadMedida, linea.unidad_medida_id)
            fila = dict(base_fila)
            fila.update(
                {
                    "Id Producto": producto.id if producto else None,
                    "Producto": producto.nombre if producto else None,
                    "Cantidad": linea.cantidad,
                    "Unidad de medida": unidad.nombre if unidad else None,
                    "Valor Material": linea.valor_material,
                    "Valor Flete": linea.valor_flete,
                    "Numero de viajes": linea.numero_viajes,
                    "Valor de comision": linea.valor_comision,
                    "Valor de costo": linea.valor_costo,
                    "Valor cama baja": linea.valor_camabaja,
                }
            )
            filas.append(fila)

    return _escribir_xlsx(filas, COLUMNAS_GENERAL)


def generar_reporte_facturadas(
    session: Session,
    *,
    desde: date | None,
    hasta: date | None,
    cliente_id: int | None,
    proyecto_id: int | None,
    vehiculo_id: int | None,
    camabaja_id: int | None,
    estado_facturacion: str | None = None,
) -> BytesIO:
    statement = _aplicar_filtros(
        select(Remision).where(Remision.estado_facturacion != "NO_FACTURADA"),
        desde=desde,
        hasta=hasta,
        cliente_id=cliente_id,
        proyecto_id=proyecto_id,
        vehiculo_id=vehiculo_id,
        camabaja_id=camabaja_id,
        estado_facturacion=estado_facturacion,
    )
    remisiones = session.exec(statement.order_by(Remision.fecha, Remision.id)).all()
    cat = _Catalogos(session)

    filas: list[dict[str, Any]] = []
    for remision in remisiones:
        activas = session.exec(
            select(RemisionFacturacion).where(
                RemisionFacturacion.remision_id == remision.id,
                RemisionFacturacion.anulada.is_(False),
            )
        ).all()
        if not activas:
            continue  # estado inconsistente puntual; no hay facturación activa que listar

        comunes = _datos_comunes(cat, remision)
        totales = _totales_remision(session, remision.id)
        valor_facturado_acumulado = sum((r.valor_facturado for r in activas), Decimal(0))
        valor_pendiente = totales["total"] - valor_facturado_acumulado

        for registro in activas:
            factura = cat.get(Factura, registro.factura_id)
            usuario_facturacion = cat.get(Usuario, registro.usuario_id)
            filas.append(
                {
                    "Numero de remision": comunes["numero_remision"],
                    "Fecha": comunes["fecha"],
                    "Forma de pago": comunes["forma_pago"],
                    "ODS": comunes["ods"],
                    "ID cliente": comunes["cliente_identificacion"],
                    "Cliente": comunes["cliente_nombre"],
                    "Proyecto": comunes["proyecto"],
                    "Tramo": comunes["tramo"],
                    "Origen": comunes["origen"],
                    "Destino": comunes["destino"],
                    "Kilometros": comunes["kilometros"],
                    "Vehiculos": comunes["vehiculo_placa"],
                    "Cama Baja": comunes["camabaja_placa"],
                    "Hora llegada": comunes["hora_llegada"],
                    "Hola salida": comunes["hora_salida"],
                    "Conductor": comunes["conductor_nombre"],
                    "Tipo de transporte": comunes["tipo_servicio"],
                    "Usuario": comunes["usuario_creador_nombre"],
                    "ID usuario": comunes["usuario_creador_id"],
                    "ID factura": factura.id if factura else None,
                    "Usuario Factura": usuario_facturacion.nombre if usuario_facturacion else None,
                    "ID usuario factura": usuario_facturacion.id if usuario_facturacion else None,
                    "Número de factura": factura.numero_factura if factura else None,
                    "Fecha de factura": factura.fecha_factura if factura else None,
                    "Valor flete": totales["valor_flete"],
                    "Valor Camabaja": totales["valor_camabaja"],
                    "Subtotal": totales["subtotal"],
                    "Valor Material": totales["valor_material"],
                    "Total": totales["total"],
                    "Valor facturado": registro.valor_facturado,
                    "Valor pendiente": valor_pendiente,
                    "Estado": comunes["estado_facturacion"],
                }
            )

    return _escribir_xlsx(filas, COLUMNAS_FACTURADAS)
