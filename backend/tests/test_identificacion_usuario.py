"""Identificación (cédula) opcional del usuario: se edita en Administración
→ Usuarios y sale en "ID del usuario" / "ID usuario" / "ID usuario factura"
de los reportes; vacía si el usuario no la tiene (no se inventa).

API sobre la base SQLite de conftest.py; reportes sobre el clon
(TEST_DATABASE_URL) dentro de una transacción que se revierte."""

import openpyxl
import pytest
from sqlmodel import Session, select

from app.models.facturacion import RemisionFacturacion
from app.models.remision import Remision
from app.models.usuario import Permiso, Rol, RolPermiso, Usuario, UsuarioRol
from app.services import reportes_service as svc
from tests.conftest import PASSWORD_PRUEBA

FILTROS = {
    "desde": None, "hasta": None, "cliente_id": None, "proyecto_id": None,
    "vehiculo_id": None, "camabaja_id": None,
}


# ── API (Administración → Usuarios) ──────────────────────────────────────────


@pytest.fixture
def admin(client_api, engine_prueba):
    with Session(engine_prueba) as s:
        usuario = s.exec(select(Usuario)).one()
        rol = Rol(nombre="Administrador")
        permiso = Permiso(codigo="usuarios.administrar", modulo="admin")
        s.add_all([rol, permiso])
        s.flush()
        s.add_all(
            [RolPermiso(rol_id=rol.id, permiso_id=permiso.id), UsuarioRol(usuario_id=usuario.id, rol_id=rol.id)]
        )
        s.commit()
    token = client_api.post(
        "/api/auth/login", data={"username": "prueba@example.com", "password": PASSWORD_PRUEBA}
    ).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def _crear(client, admin, **extra):
    return client.post(
        "/api/admin/usuarios",
        json={"nombre": "Ana", "email": "ana@example.com", "password": "ClaveAna123", **extra},
        headers=admin,
    )


def test_crear_sin_identificacion_queda_vacia(client_api, admin):
    r = _crear(client_api, admin)
    assert r.status_code == 201
    assert r.json()["identificacion"] is None


def test_crear_con_identificacion_conserva_ceros_y_recorta_espacios(client_api, admin):
    r = _crear(client_api, admin, identificacion="  0012345678 ")
    assert r.status_code == 201
    assert r.json()["identificacion"] == "0012345678"


def test_editar_y_borrar_identificacion(client_api, admin):
    uid = _crear(client_api, admin).json()["id"]
    r = client_api.put(f"/api/admin/usuarios/{uid}", json={"identificacion": "1085275375"}, headers=admin)
    assert r.json()["identificacion"] == "1085275375"
    # Editar otro campo no la toca.
    r = client_api.put(f"/api/admin/usuarios/{uid}", json={"nombre": "Ana María"}, headers=admin)
    assert r.json()["identificacion"] == "1085275375"
    # Vaciarla la deja sin dato.
    r = client_api.put(f"/api/admin/usuarios/{uid}", json={"identificacion": "   "}, headers=admin)
    assert r.json()["identificacion"] is None
    listado = client_api.get("/api/admin/usuarios", headers=admin).json()["items"]
    assert {u["id"]: u["identificacion"] for u in listado}[uid] is None


@pytest.mark.parametrize("mala", ["12 345", "123;DROP", "x" * 31, "12345$"])
def test_identificacion_invalida_se_rechaza(client_api, admin, mala):
    r = _crear(client_api, admin, identificacion=mala)
    assert r.status_code == 422


def test_editar_identificacion_no_cierra_sesiones(client_api, admin, engine_prueba):
    uid = _crear(client_api, admin).json()["id"]
    client_api.put(f"/api/admin/usuarios/{uid}", json={"identificacion": "123"}, headers=admin)
    with Session(engine_prueba) as s:
        assert s.get(Usuario, uid).token_version == 0


# ── Reportes ────────────────────────────────────────────────────────────────


@pytest.fixture
def session_con_rollback(pg_engine):
    conexion = pg_engine.connect()
    transaccion = conexion.begin()
    session = Session(bind=conexion, join_transaction_mode="create_savepoint")
    try:
        yield session
    finally:
        session.close()
        transaccion.rollback()
        conexion.close()


def _filas(buffer):
    ws = openpyxl.load_workbook(buffer).active
    encabezados = [c.value for c in ws[1]]
    return [dict(zip(encabezados, (c.value for c in fila), strict=True)) for fila in ws.iter_rows(min_row=2)]


def test_reportes_vacios_si_el_usuario_no_tiene_identificacion(session_con_rollback):
    s = session_con_rollback
    for u in s.exec(select(Usuario)).all():
        u.identificacion = None
        s.add(u)
    s.flush()
    general = _filas(svc.generar_reporte_general(s, **FILTROS))
    facturadas = _filas(svc.generar_reporte_facturadas(s, **FILTROS))
    assert all(f["ID del usuario"] is None for f in general)
    assert all(f["ID usuario"] is None and f["ID usuario factura"] is None for f in facturadas)
    # Ya no sale el id interno (antes "1").
    assert all(f["ID del usuario"] != "1" for f in general)


def test_reportes_usan_la_identificacion_cuando_existe(session_con_rollback):
    s = session_con_rollback
    creador_id = s.exec(select(Remision.usuario_creador_id)).first()
    facturador_id = s.exec(select(RemisionFacturacion.usuario_id)).first()
    creador = s.get(Usuario, creador_id)
    creador.identificacion = "0087491991"
    s.add(creador)
    s.flush()
    remisiones_del_creador = {
        r.numero_remision for r in s.exec(select(Remision).where(Remision.usuario_creador_id == creador_id)).all()
    }

    general = _filas(svc.generar_reporte_general(s, **FILTROS))
    for f in general:
        esperado = "0087491991" if f["Numero de remision"] in remisiones_del_creador else None
        assert f["ID del usuario"] == esperado, f["Numero de remision"]

    facturadas = _filas(svc.generar_reporte_facturadas(s, **FILTROS))
    for f in facturadas:
        if f["Numero de remision"] in remisiones_del_creador:
            assert f["ID usuario"] == "0087491991"
    if facturador_id == creador_id:
        assert all(f["ID usuario factura"] == "0087491991" for f in facturadas)


def test_no_se_reasigna_el_creador_de_las_historicas(session_con_rollback):
    # La identificación solo se lee: el usuario_creador_id de las remisiones
    # no cambia por generar reportes ni por editar usuarios.
    s = session_con_rollback
    antes = dict(s.exec(select(Remision.id, Remision.usuario_creador_id)).all())
    svc.generar_reporte_general(s, **FILTROS)
    assert dict(s.exec(select(Remision.id, Remision.usuario_creador_id)).all()) == antes
