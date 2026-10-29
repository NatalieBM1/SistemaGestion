"""Verifica que los reportes de Remisiones salgan con el formato exacto que
pidió el cliente: mismos encabezados y mismo orden que la columna E de los
Excel de referencia en ``app/data_import/formatos/``, y que el reporte
general traiga una fila por línea de producto.

Corre contra la base de TEST_DATABASE_URL (solo lectura, no escribe
nada; ver conftest.pg_engine) — pensado para correrse contra una base con datos reales o de
ensayo, no requiere una base vacía.
"""

from pathlib import Path

import openpyxl
import pytest
from sqlmodel import Session, func, select

from app.models.remision import Remision, RemisionProducto
from app.services import reportes_service as svc

FORMATOS_DIR = Path(__file__).resolve().parents[1] / "app" / "data_import" / "formatos"


def _encabezados_de_referencia(nombre_archivo: str) -> list[str]:
    """Columna E (fila 7 en adelante) del Excel de referencia, tal cual —
    incluidos los errores del original (ej. "Hola salida"): el cliente
    pidió expresamente conservarlos."""
    ruta = FORMATOS_DIR / nombre_archivo
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb.active
    encabezados = []
    for row in range(7, ws.max_row + 1):
        valor = ws.cell(row=row, column=5).value
        if valor is None:
            break
        encabezados.append(valor)
    return encabezados


def _encabezados_generados(buffer) -> list[str]:
    wb = openpyxl.load_workbook(buffer)
    return [celda.value for celda in wb.active[1]]


@pytest.fixture(scope="module")
def session(pg_engine):
    with Session(pg_engine) as s:
        yield s


def _assert_encabezados_identicos(generados: list[str], esperados: list[str]) -> None:
    """Compara carácter por carácter (== de strings ya lo hace) y, si
    falla, señala exactamente en qué posición y columna difiere, en vez
    de solo decir "las listas no son iguales"."""
    assert len(generados) == len(esperados), (
        f"Se esperaban {len(esperados)} columnas, se generaron {len(generados)}."
    )
    diferencias = [
        f"  columna {i + 1}: generado={g!r} vs esperado={e!r}"
        for i, (g, e) in enumerate(zip(generados, esperados))
        if g != e
    ]
    assert not diferencias, "Encabezados distintos:\n" + "\n".join(diferencias)


def test_encabezados_reporte_general(session):
    buffer = svc.generar_reporte_general(
        session,
        desde=None,
        hasta=None,
        cliente_id=None,
        proyecto_id=None,
        vehiculo_id=None,
        camabaja_id=None,
    )
    generados = _encabezados_generados(buffer)
    esperados = _encabezados_de_referencia("reporte-remisiones-general-programa_nuevo.xlsx")
    _assert_encabezados_identicos(generados, esperados)


def test_encabezados_reporte_facturadas(session):
    buffer = svc.generar_reporte_facturadas(
        session,
        desde=None,
        hasta=None,
        cliente_id=None,
        proyecto_id=None,
        vehiculo_id=None,
        camabaja_id=None,
    )
    generados = _encabezados_generados(buffer)
    esperados = _encabezados_de_referencia("reporte-remisiones-facturadas-programa_nuevo.xlsx")
    _assert_encabezados_identicos(generados, esperados)


def test_reporte_general_una_fila_por_linea_de_producto(session):
    """El reporte general trae una fila por línea de producto; una remisión
    sin ninguna línea cuenta como una sola fila (sus columnas de producto
    quedan vacías)."""
    buffer = svc.generar_reporte_general(
        session,
        desde=None,
        hasta=None,
        cliente_id=None,
        proyecto_id=None,
        vehiculo_id=None,
        camabaja_id=None,
    )
    wb = openpyxl.load_workbook(buffer)
    filas_generadas = wb.active.max_row - 1  # sin contar el encabezado

    total_lineas = session.exec(select(func.count()).select_from(RemisionProducto)).one()
    total_remisiones = session.exec(select(func.count()).select_from(Remision)).one()
    remisiones_con_lineas = session.exec(
        select(func.count(func.distinct(RemisionProducto.remision_id)))
    ).one()
    remisiones_sin_ninguna_linea = total_remisiones - remisiones_con_lineas

    assert filas_generadas == total_lineas + remisiones_sin_ninguna_linea


# --- tipo y formato de dato por columna ---
# Verifica las reglas que pidió el cliente: fechas como fecha (formato de
# celda yyyy-mm-dd), identificaciones/códigos como texto (para no perder
# ceros a la izquierda ni caer en notación científica), y valores
# monetarios como número plano. Los archivos de referencia en formatos/
# solo traen los encabezados (columna E), no datos de ejemplo en el
# formato nuevo — así que esto valida la salida propia contra las reglas
# del pedido, no contra una celda de muestra del Excel del cliente.

COLUMNAS_TEXTO_ESPERADAS = {
    "ID del cliente", "ID cliente", "ID del conductor",
    "ID del usuario", "ID usuario", "ID usuario factura",
}
COLUMNAS_FECHA_ESPERADAS = {"Fecha", "Fecha de factura"}
COLUMNAS_DINERO_ESPERADAS_GENERAL = {
    "Valor Material", "Valor Flete", "Valor de comision", "Valor de costo", "Valor cama baja",
}
COLUMNAS_DINERO_ESPERADAS_FACTURADAS = {
    "Valor flete", "Valor Camabaja", "Subtotal", "Valor Material", "Total",
    "Valor facturado", "Valor pendiente",
}


def _verificar_tipos_por_columna(buffer, columnas_texto, columnas_fecha, columnas_dinero):
    import datetime
    from decimal import Decimal

    wb = openpyxl.load_workbook(buffer)
    ws = wb.active
    encabezados = [c.value for c in ws[1]]
    indice = {nombre: i + 1 for i, nombre in enumerate(encabezados)}

    for fila in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 50)):
        for nombre, col in indice.items():
            celda = ws.cell(row=fila[0].row, column=col)
            if celda.value is None:
                continue  # celda vacía por falta de dato: válido, no se verifica tipo
            if nombre in columnas_texto:
                assert isinstance(celda.value, str), (
                    f"{nombre} (fila {celda.row}) debería ser texto, es {type(celda.value)}: {celda.value!r}"
                )
            elif nombre in columnas_fecha:
                assert isinstance(celda.value, (datetime.date, datetime.datetime)), (
                    f"{nombre} (fila {celda.row}) debería ser fecha, es {type(celda.value)}: {celda.value!r}"
                )
                assert "yyyy-mm-dd" in celda.number_format.lower() or celda.number_format == "yyyy-mm-dd", (
                    f"{nombre} (fila {celda.row}) tiene formato de celda {celda.number_format!r}, se esperaba yyyy-mm-dd"
                )
            elif nombre in columnas_dinero:
                assert isinstance(celda.value, (int, float, Decimal)) and not isinstance(celda.value, bool), (
                    f"{nombre} (fila {celda.row}) debería ser número, es {type(celda.value)}: {celda.value!r}"
                )


def test_tipos_de_dato_reporte_general(session):
    buffer = svc.generar_reporte_general(
        session, desde=None, hasta=None, cliente_id=None, proyecto_id=None,
        vehiculo_id=None, camabaja_id=None,
    )
    _verificar_tipos_por_columna(
        buffer, COLUMNAS_TEXTO_ESPERADAS, COLUMNAS_FECHA_ESPERADAS, COLUMNAS_DINERO_ESPERADAS_GENERAL
    )


def test_tipos_de_dato_reporte_facturadas(session):
    buffer = svc.generar_reporte_facturadas(
        session, desde=None, hasta=None, cliente_id=None, proyecto_id=None,
        vehiculo_id=None, camabaja_id=None,
    )
    _verificar_tipos_por_columna(
        buffer, COLUMNAS_TEXTO_ESPERADAS, COLUMNAS_FECHA_ESPERADAS, COLUMNAS_DINERO_ESPERADAS_FACTURADAS
    )


# --- lista exacta pedida por el cliente (Bloque C) ---
# Escrita a mano a propósito, separada de la hoja de referencia y del
# código: si alguien cambia un nombre, el orden o una tilde en cualquiera
# de los tres lados, falla.

ENCABEZADOS_GENERAL_CLIENTE = (
    "ID | Numero de remision | Fecha | Forma de pago | ODS | ID del cliente | Cliente | "
    "Proyecto | Tramo | Origen | Destino | Kilometros | Vehiculos | Cama Baja | "
    "Hora llegada | Hola salida | ID del conductor | Conductor | Tipo de transporte | "
    "Observaciones | Id Producto | Producto | Cantidad | Unidad de medida | Valor Material | "
    "Valor Flete | Numero de viajes | Valor de comision | Valor de costo | Valor cama baja | "
    "Remision AGRESUR | ID del usuario | Usuario | Estado"
).split(" | ")

ENCABEZADOS_FACTURADAS_CLIENTE = (
    "Numero de remision | Fecha | Forma de pago | ODS | ID cliente | Cliente | Proyecto | "
    "Tramo | Origen | Destino | Kilometros | Vehiculos | Cama Baja | Hora llegada | "
    "Hola salida | Conductor | Tipo de transporte | Usuario | ID usuario | ID factura | "
    "Usuario Factura | ID usuario factura | Número de factura | Fecha de factura | "
    "Valor flete | Valor Camabaja | Subtotal | Valor Material | Total | Valor facturado | "
    "Valor pendiente | Estado"
).split(" | ")

FILTROS = {
    "desde": None, "hasta": None, "cliente_id": None, "proyecto_id": None,
    "vehiculo_id": None, "camabaja_id": None,
}


@pytest.fixture(scope="module")
def hoja_general(session):
    return openpyxl.load_workbook(svc.generar_reporte_general(session, **FILTROS))


@pytest.fixture(scope="module")
def hoja_facturadas(session):
    return openpyxl.load_workbook(svc.generar_reporte_facturadas(session, **FILTROS))


def test_cantidad_de_columnas_pedida():
    assert len(ENCABEZADOS_GENERAL_CLIENTE) == 34
    assert len(ENCABEZADOS_FACTURADAS_CLIENTE) == 32


@pytest.mark.parametrize(
    ("fixture", "esperados", "referencia"),
    [
        ("hoja_general", ENCABEZADOS_GENERAL_CLIENTE, "reporte-remisiones-general-programa_nuevo.xlsx"),
        ("hoja_facturadas", ENCABEZADOS_FACTURADAS_CLIENTE, "reporte-remisiones-facturadas-programa_nuevo.xlsx"),
    ],
)
def test_encabezados_identicos_a_la_lista_del_cliente(request, fixture, esperados, referencia):
    wb = request.getfixturevalue(fixture)
    generados = [c.value for c in wb.active[1]]
    _assert_encabezados_identicos(generados, esperados)
    # Y la hoja de referencia dice lo mismo que la lista.
    _assert_encabezados_identicos(_encabezados_de_referencia(referencia), esperados)


@pytest.mark.parametrize("fixture", ["hoja_general", "hoja_facturadas"])
def test_importable_directo_fila_1_sin_combinar_ni_titulo(request, fixture):
    wb = request.getfixturevalue(fixture)
    ws = wb.active
    assert len(wb.worksheets) == 1
    assert not ws.merged_cells.ranges
    assert all(isinstance(c.value, str) and c.value for c in ws[1])
    # La fila 2 ya es un dato (no hay fila de encabezados cortos).
    assert ws.cell(row=2, column=1).value not in (None, "id", "numero_remision")


def _columna(ws, nombre):
    encabezados = [c.value for c in ws[1]]
    col = encabezados.index(nombre) + 1
    return [ws.cell(row=r, column=col) for r in range(2, ws.max_row + 1)]


@pytest.mark.parametrize(
    ("fixture", "columnas"),
    [
        ("hoja_general", ["ID del cliente", "ID del conductor", "ID del usuario"]),
        ("hoja_facturadas", ["ID cliente", "ID usuario", "ID usuario factura"]),
    ],
)
def test_identificaciones_como_texto_en_todas_las_filas(request, fixture, columnas):
    ws = request.getfixturevalue(fixture).active
    for nombre in columnas:
        for celda in _columna(ws, nombre):
            if celda.value is None:
                continue
            assert isinstance(celda.value, str) and celda.number_format == "@", (nombre, celda.coordinate)
            assert "E+" not in celda.value.upper(), (nombre, celda.value)


@pytest.mark.parametrize("fixture", ["hoja_general", "hoja_facturadas"])
def test_horas_con_formato_hh_mm(request, fixture):
    import datetime

    ws = request.getfixturevalue(fixture).active
    vistas = 0
    for nombre in ("Hora llegada", "Hola salida"):
        for celda in _columna(ws, nombre):
            if celda.value is None:
                continue
            vistas += 1
            assert isinstance(celda.value, datetime.time), (nombre, celda.coordinate, celda.value)
            assert celda.number_format == "hh:mm", (nombre, celda.number_format)
    assert vistas > 0


def test_dinero_como_numero_entero_si_no_tiene_centavos(hoja_general, session):
    ws = hoja_general.active
    for nombre in ("Valor Material", "Valor Flete", "Valor de comision", "Valor de costo", "Valor cama baja"):
        for celda in _columna(ws, nombre):
            if celda.value is None:
                continue
            assert isinstance(celda.value, (int, float)) and not isinstance(celda.value, bool)
            if isinstance(celda.value, float):
                assert celda.value != int(celda.value), "un entero no debe ir como 250000.0"
    # Los centavos que sí existen en la base se conservan.
    lineas = session.exec(select(RemisionProducto)).all()
    with_cents = {float(p.valor_flete) for p in lineas if p.valor_flete % 1}
    en_reporte = {c.value for c in _columna(ws, "Valor Flete") if isinstance(c.value, float)}
    assert with_cents <= en_reporte


def test_numero_de_remision_con_el_formato_guardado(hoja_general, session):
    guardados = set(session.exec(select(Remision.numero_remision)).all())
    en_reporte = {c.value for c in _columna(hoja_general.active, "Numero de remision")}
    assert en_reporte == guardados


def test_facturadas_una_fila_por_remision_y_factura(hoja_facturadas, session):
    from app.models.facturacion import RemisionFacturacion

    activas = session.exec(
        select(func.count()).select_from(RemisionFacturacion).where(RemisionFacturacion.anulada.is_(False))
    ).one()
    assert hoja_facturadas.active.max_row - 1 == activas
